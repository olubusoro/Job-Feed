"""
Excel export service using openpyxl.
Generates an .xlsx file with formatted columns, clickable hyperlinks,
frozen header row, and auto-filter.
"""

from __future__ import annotations

import io
from datetime import datetime
from typing import Optional

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models import Job, UserJob


def _time_ago(dt: Optional[datetime]) -> str:
    """Return a human-readable relative time string."""
    if dt is None:
        return "Unknown"
    delta = datetime.utcnow() - dt
    hours = int(delta.total_seconds() // 3600)
    if hours < 1:
        return "Just now"
    if hours < 24:
        return f"{hours}h ago"
    days = hours // 24
    return f"{days}d ago"


def generate_excel(user_jobs: list[UserJob]) -> bytes:
    """
    Generate an .xlsx file from a list of UserJob rows.
    Returns raw bytes suitable for an HTTP response.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "JobFeed Results"

    # ── Header row ────────────────────────────────────────────────────────────
    headers = [
        "Rank",
        "Company",
        "Job Title",
        "Location",
        "Posted",
        "Score",
        "ATS",
        "Apply Link",
        "Status",
        "Notes",
    ]

    header_fill = PatternFill("solid", fgColor="1A1A2E")
    header_font = Font(bold=True, color="E94560", size=11)

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Freeze header
    ws.freeze_panes = "A2"

    # Autofilter
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    # ── Data rows ────────────────────────────────────────────────────────────
    link_font = Font(color="0563C1", underline="single")

    for row_idx, uj in enumerate(user_jobs, start=2):
        job: Job = uj.job
        posted_str = _time_ago(job.posted_at or job.first_seen_at)

        row_data = [
            uj.rank,
            job.company.name if job.company else "",
            job.title,
            job.location,
            posted_str,
            uj.score,
            job.ats,
            None,          # Apply Link — handled separately as hyperlink
            uj.status,
            uj.notes or "",
        ]

        for col_idx, value in enumerate(row_data, start=1):
            if value is None:
                continue
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            # Alternate row shading
            if row_idx % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F5F5F5")

        # Clickable hyperlink in "Apply Link" column (col 8)
        if job.url:
            link_cell = ws.cell(row=row_idx, column=8, value="Apply →")
            link_cell.hyperlink = job.url
            link_cell.font = link_font
            link_cell.alignment = Alignment(horizontal="center")

    # ── Column widths ─────────────────────────────────────────────────────────
    col_widths = [6, 24, 36, 22, 12, 8, 14, 12, 12, 30]
    for col_idx, width in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Row height for header
    ws.row_dimensions[1].height = 22

    # ── Save to bytes ────────────────────────────────────────────────────────
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()
